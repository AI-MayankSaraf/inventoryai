"""
Turning an uploaded file into a table, without a model.

Spreadsheets and CSVs already *are* tables; the only real problem is that
a supplier's quotation rarely starts at row 1. There is usually a letterhead,
an address block, maybe a subject line, and then the grid. So the parser's
job is to find the header row and read from there.

§5 states the rule as "the first row with at least three non-empty,
non-numeric cells", and that rule alone is wrong for CSVs: an address line
like `Plot 14, MIDC Industrial Area, Pune 411019` is three non-numeric
cells, so the letterhead wins and the real table is read as data. The fix
is to let the caller pass the words it already knows are column headings —
the `canonical_fields` synonym vocabulary — and prefer the row that
contains most of them. The §5 rule stays as the fallback for a layout whose
headings we have never seen.

PDFs are read for text and for tables. A PDF with no extractable text is a
scan, and a scan needs OCR, which needs a provider — that comes back as a
clear failure rather than an empty extraction that looks like a blank
document.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from typing import Optional

from app.modules.ai import providers
from app.modules.ai.normalise import normalise

#: BR-DOC-01, verified against the sniffed content below rather than trusted.
SUPPORTED_EXTENSIONS = {"xlsx", "xls", "csv", "pdf", "docx", "jpg", "jpeg", "png", "webp"}

IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "webp"}

#: What a file actually is, by its first bytes. An `.xlsx` that is really an
#: executable, or a `.pdf` that is really a zip, never reaches a parser.
_MAGIC: list[tuple[bytes, set[str]]] = [
    (b"PK\x03\x04", {"xlsx", "docx"}),            # any OOXML/zip container
    (b"%PDF-", {"pdf"}),
    (b"\xff\xd8\xff", {"jpg", "jpeg"}),
    (b"\x89PNG\r\n\x1a\n", {"png"}),
    (b"\xd0\xcf\x11\xe0", {"xls"}),               # legacy OLE2 compound file
]


class UnsupportedFile(ValueError):
    def __init__(self, message: str, code: str = "UNSUPPORTED_FILE") -> None:
        super().__init__(message)
        self.code = code


@dataclass
class ParsedTable:
    sheet_name: str
    header_row_index: int
    data_start_row: int
    headers: list[str]
    rows: list[list[str]]

    @property
    def column_signature(self) -> str:
        """`sha256` of the normalised, ordered header texts (§5 step 2).

        Normalised, so "Rate " and "rate" are the same layout; ordered, so
        two suppliers using the same words in a different order are not.
        """
        joined = "|".join(normalise(h) for h in self.headers)
        return hashlib.sha256(joined.encode()).hexdigest()


@dataclass
class ParsedDocument:
    file_format: str
    text: str = ""
    page_count: Optional[int] = None
    tables: list[ParsedTable] = field(default_factory=list)
    #: Notes for the pipeline trace — what the parser did and did not find.
    notes: list[str] = field(default_factory=list)

    @property
    def primary(self) -> Optional[ParsedTable]:
        """The table with the most data rows. A workbook's first sheet is
        often a cover page, so "first" is the wrong heuristic."""
        return max(self.tables, key=lambda t: len(t.rows), default=None)


def sniff(blob: bytes, extension: str) -> str:
    """Return the extension the *content* supports, or raise.

    A zip container could be either `xlsx` or `docx` and the extension is
    the only thing that distinguishes them, so within a magic group the
    claimed extension is accepted; across groups it is not.
    """
    ext = extension.lower().lstrip(".")
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFile(
            f".{ext} files are not supported. Upload a spreadsheet, CSV, PDF, Word file or photo."
        )
    image = _image_format(blob)
    if image and ext in IMAGE_EXTENSIONS:
        # "Save image as" routinely writes a WebP or JPEG under a .png name.
        # Every image goes to the same OCR path, so the extension only has
        # to say "image"; the bytes decide which kind.
        return image
    for magic, group in _MAGIC:
        if blob.startswith(magic):
            if ext in group:
                return ext
            raise UnsupportedFile(
                f"This file is named .{ext} but its contents are not. Re-save it and try again.",
                code="FILE_CONTENT_MISMATCH",
            )
    if ext == "csv":
        # CSV has no magic number; it just has to decode as text.
        try:
            blob[:4096].decode("utf-8-sig")
        except UnicodeDecodeError:
            raise UnsupportedFile("This file is named .csv but is not text.", code="FILE_CONTENT_MISMATCH")
        return "csv"
    raise UnsupportedFile("This file type could not be recognised.", code="FILE_CONTENT_MISMATCH")


def _image_format(blob: bytes) -> Optional[str]:
    if blob.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if blob.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if blob[:4] == b"RIFF" and blob[8:12] == b"WEBP":
        return "webp"
    return None


def parse(blob: bytes, extension: str, hints: Optional[set[str]] = None) -> ParsedDocument:
    """`hints` are normalised words known to be column headings — pass the
    canonical vocabulary's synonyms so the header row is found by meaning
    rather than by shape."""
    ext = sniff(blob, extension)
    if ext == "csv":
        return _parse_csv(blob, hints)
    if ext == "xlsx":
        return _parse_xlsx(blob, hints)
    if ext == "pdf":
        return _parse_pdf(blob, hints)
    if ext in IMAGE_EXTENSIONS:
        return _parse_image(blob, ext, hints)
    if ext == "docx":
        return _parse_docx(blob, hints)
    raise UnsupportedFile(
        f".{ext} cannot be read by this installation. Save it as .xlsx or .csv and upload that.",
    )


# ------------------------------------------------------------------ tables


def _is_headerish(cell: str) -> bool:
    """A header cell has words in it, not a number and not emptiness."""
    text = (cell or "").strip()
    if not text or len(text) > 60:
        return False
    return not re.fullmatch(r"[\d\s.,%₹/-]+", text)


def _hinted_header_row(grid: list[list[str]], hints: Optional[set[str]], limit: int = 25) -> Optional[int]:
    """The row containing the most known column headings, if any row has
    at least two. `None` when the vocabulary recognises nothing."""
    if not hints:
        return None
    best_index, best_score = None, 0
    for index, row in enumerate(grid[:limit]):
        score = sum(1 for cell in row if normalise(cell) in hints)
        # Two recognised headings is already far better evidence than
        # "three cells with letters in them"; one could be a coincidence.
        if score >= 2 and score > best_score:
            best_index, best_score = index, score
    return best_index


def _find_header_row(grid: list[list[str]], hints: Optional[set[str]] = None) -> int:
    """Which row holds the column headings.

    Scanning only the first 25 rows on purpose — a heading further down
    than that is not a letterhead, it is a different document.
    """
    hinted = _hinted_header_row(grid, hints)
    if hinted is not None:
        return hinted

    for index, row in enumerate(grid[:25]):
        if sum(1 for cell in row if _is_headerish(cell)) >= 3:
            return index
    return 0


def _table_from_grid(grid: list[list[str]], sheet_name: str, hints: Optional[set[str]] = None) -> Optional[ParsedTable]:
    grid = [row for row in grid]
    if not grid:
        return None
    header_index = _find_header_row(grid, hints)
    headers = [(cell or "").strip() for cell in grid[header_index]]
    # Trailing empties are spreadsheet padding, not columns.
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        return None
    # Two columns both headed "Amount" (or, from OCR, both read as the same
    # smudge) are still two columns. The mapping is keyed by heading text,
    # so the second becomes "Amount (2)" rather than colliding with the first.
    seen: dict[str, int] = {}
    for index, header in enumerate(headers):
        key = normalise(header)
        seen[key] = seen.get(key, 0) + 1
        if seen[key] > 1:
            headers[index] = f"{header} ({seen[key]})"
    width = len(headers)

    rows: list[list[str]] = []
    for row in grid[header_index + 1 :]:
        cells = [(cell or "").strip() for cell in row[:width]]
        cells += [""] * (width - len(cells))
        if any(cells):
            rows.append(cells)
    return ParsedTable(
        sheet_name=sheet_name,
        header_row_index=header_index,
        data_start_row=header_index + 1,
        headers=headers,
        rows=rows,
    )


def _parse_csv(blob: bytes, hints: Optional[set[str]] = None) -> ParsedDocument:
    text = blob.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    grid = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    doc = ParsedDocument(file_format="csv", text=text[:200_000], page_count=1)
    table = _table_from_grid(grid, "csv", hints)
    if table:
        doc.tables.append(table)
        doc.notes.append(f"Read {len(table.rows)} rows under a header on line {table.header_row_index + 1}.")
    else:
        doc.notes.append("The file had no readable rows.")
    return doc


def _parse_xlsx(blob: bytes, hints: Optional[set[str]] = None) -> ParsedDocument:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    doc = ParsedDocument(file_format="xlsx", page_count=len(workbook.sheetnames))
    text_parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            grid: list[list[str]] = []
            for row in sheet.iter_rows(max_row=2000, values_only=True):
                grid.append(["" if cell is None else str(cell) for cell in row])
            text_parts.append("\n".join("\t".join(row) for row in grid[:200]))
            table = _table_from_grid(grid, sheet.title, hints)
            if table and table.rows:
                doc.tables.append(table)
    finally:
        workbook.close()
    doc.text = "\n".join(text_parts)[:200_000]
    if doc.tables:
        best = doc.primary
        assert best is not None
        doc.notes.append(
            f"Read {len(best.rows)} rows from sheet \"{best.sheet_name}\" "
            f"under a header on row {best.header_row_index + 1}."
        )
    else:
        doc.notes.append("No sheet in the workbook had a readable table.")
    return doc


def _parse_pdf(blob: bytes, hints: Optional[set[str]] = None) -> ParsedDocument:
    import pdfplumber

    doc = ParsedDocument(file_format="pdf_table")
    text_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(blob)) as pdf:
        doc.page_count = len(pdf.pages)
        for page_number, page in enumerate(pdf.pages[:20], start=1):
            text_parts.append(page.extract_text() or "")
            for raw in page.extract_tables() or []:
                grid = [["" if cell is None else str(cell) for cell in row] for row in raw]
                table = _table_from_grid(grid, f"page {page_number}", hints)
                if table and table.rows:
                    doc.tables.append(table)
    doc.text = "\n".join(text_parts)[:200_000]

    if not doc.text.strip():
        # No text layer at all: this is a scan, and reading it is OCR's job.
        if not providers.ocr.configured:
            raise UnsupportedFile(
                "This PDF is a scan with no text in it, and no OCR provider is configured. "
                "Upload the spreadsheet or a text PDF instead.",
                code="OCR_NOT_CONFIGURED",
            )
        doc.text, grids = providers.ocr.read_document(blob, "application/pdf", hints)
        for name, grid in grids:
            table = _table_from_grid(grid, name, hints)
            if table and table.rows:
                doc.tables.append(table)
        doc.notes.append("This PDF is a scan; it was read by OCR. Check the values against the page.")
    if doc.tables:
        best = doc.primary
        assert best is not None
        doc.notes.append(f"Found a {len(best.headers)}-column table on {best.sheet_name}.")
    else:
        doc.notes.append("No table was found; only the header fields could be read from the text.")
    return doc


def _parse_image(blob: bytes, ext: str, hints: Optional[set[str]] = None) -> ParsedDocument:
    if not providers.ocr.configured:
        raise UnsupportedFile(
            "Photos and scans need an OCR provider, and none is configured on this installation. "
            "Upload the supplier's spreadsheet or PDF instead.",
            code="OCR_NOT_CONFIGURED",
        )
    text, grids = providers.ocr.read_document(blob, f"image/{'jpeg' if ext in {'jpg', 'jpeg'} else ext}", hints)
    doc = ParsedDocument(file_format="pdf_table", text=text, page_count=1,
                         notes=["This is an image; it was read by OCR. Check the values against the page."])
    for name, grid in grids:
        table = _table_from_grid(grid, name, hints)
        if table and table.rows:
            doc.tables.append(table)
    if doc.tables:
        best = doc.primary
        doc.notes.append(f"Rebuilt a {len(best.headers)}-column table with {len(best.rows)} rows from the image.")
    else:
        doc.notes.append("No table could be rebuilt from the image; only the header fields could be read.")
        width = _image_width(blob)
        if width and width < _MIN_READABLE_WIDTH:
            # The usual reason, and the one the person can act on: a page
            # saved small (a screenshot, an image from a website) has text a
            # few pixels high, and no amount of OCR recovers detail that is
            # not in the file.
            doc.notes.append(
                f"The image is only {width} pixels wide, so the text on it is too small to read. "
                f"Upload the supplier's PDF or Excel file, or a photo or scan at least "
                f"{_MIN_READABLE_WIDTH} pixels wide."
            )
    return doc


#: Below this, a full-page document's body text is under ~10 px high —
#: the size at which Tesseract stops reading it, measured on a real
#: quotation (676 px wide: 4 of 11 table values read at best).
_MIN_READABLE_WIDTH = 1200


def _image_width(blob: bytes) -> Optional[int]:
    try:
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as image:
            return image.width
    except Exception:
        return None


def _parse_docx(blob: bytes, hints: Optional[set[str]] = None) -> ParsedDocument:
    """Word files carry their tables in XML; this reads them without a
    dependency, because a supplier quotation in .docx is a table in a
    wrapper and nothing more."""
    import xml.etree.ElementTree as ET
    import zipfile

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    doc = ParsedDocument(file_format="docx", page_count=1)
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        try:
            xml = archive.read("word/document.xml")
        except KeyError:
            raise UnsupportedFile("This .docx has no document body.", code="FILE_CONTENT_MISMATCH")
    root = ET.fromstring(xml)

    paragraphs = ["".join(t.text or "" for t in p.iter(f"{{{ns['w']}}}t")) for p in root.iter(f"{{{ns['w']}}}p")]
    doc.text = "\n".join(paragraphs)[:200_000]

    for index, table in enumerate(root.iter(f"{{{ns['w']}}}tbl")):
        grid = [
            [
                "".join(t.text or "" for t in cell.iter(f"{{{ns['w']}}}t")).strip()
                for cell in row.iter(f"{{{ns['w']}}}tc")
            ]
            for row in table.iter(f"{{{ns['w']}}}tr")
        ]
        parsed = _table_from_grid(grid, f"table {index + 1}", hints)
        if parsed and parsed.rows:
            doc.tables.append(parsed)
    doc.notes.append(
        f"Read {len(doc.tables)} table(s) from the Word file." if doc.tables else "No table was found in the Word file."
    )
    return doc
