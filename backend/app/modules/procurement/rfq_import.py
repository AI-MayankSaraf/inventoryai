"""
RFQ import — reading a buyer's spreadsheet whatever shape it arrives in.

There is no standard RFQ layout in India. The same request turns up as
"Particulars / Qty. / Unit", "Name of Item / Qty / Per" (Tally), "Material
Description / Req. Qty / UoM" (SAP), with a letterhead above the table, a
serial-number column, section headings between the lines and "10 Nos" typed
into the quantity cell. This module turns all of that into RFQ lines in three
layers, each of which the next can override:

1. **Saved layout.** A layout someone already imported is recognised by the
   signature of its header row and reused exactly, whatever row it is on.
   Stored in `document_schema_mappings` (`document_type = 'rfq_import'`) —
   the same learned-layout table the AI pipeline uses, so a confirmed RFQ
   layout also shows up on the Schema Mappings screen.
2. **Auto-detection.** The header row is found anywhere in the first 30 rows
   of any sheet, and each header is scored against a synonym list written
   for Indian procurement paperwork: exact synonym, then whole-word
   containment ("Material Description"), then typo tolerance ("Discription").
3. **Manual mapping.** The import screen shows what was detected and lets a
   person point any field at any column, or pick the header row, before
   anything is written. Importing confirms that mapping and saves it.

Nothing here writes except `save_layout`, which only runs from a real import.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from difflib import SequenceMatcher
from typing import Optional
from uuid import UUID, uuid4

from fastapi import status
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import CODE_VALIDATION, ApiError
from app.modules.ai import column_assist

log = logging.getLogger(__name__)

DOCUMENT_TYPE = "rfq_import"

#: Scan this far down for the header row. A table that starts lower than
#: this is not under a letterhead, it is in a different document.
_HEADER_SCAN_ROWS = 30
_MAX_ROWS = 5000

# ------------------------------------------------------------ vocabulary

#: field -> (label, required, synonyms). Synonyms are compared after
#: `header_key()`, so punctuation, case, "*" and bracketed notes are gone.
FIELDS: dict[str, tuple[str, bool, tuple[str, ...]]] = {
    "description": (
        "Description",
        True,
        (
            "description", "item", "items", "item description", "item desc", "desc", "product",
            "product name", "product description", "item name", "name of item", "name of the item",
            "particulars", "particular", "item particulars", "details", "item details",
            "material", "material description", "material name", "material desc",
            "specification", "specifications", "item specification", "technical specification",
            "goods", "description of goods", "goods description", "name of goods",
            "article", "article description", "commodity", "name", "item and specification",
            "description of item", "description of items", "description of material",
            "item name and description", "item description and specification",
        ),
    ),
    "quantity": (
        "Quantity",
        True,
        (
            "qty", "quantity", "req qty", "reqd qty", "required qty", "required quantity",
            "qty required", "quantity required", "qty reqd", "order qty", "order quantity",
            "total qty", "total quantity", "requirement", "reqd quantity", "no of units",
            "number of units", "units required", "est qty", "estimated qty", "estimated quantity",
            "tender qty", "enquiry qty", "rfq qty", "indent qty", "qty to be supplied",
        ),
    ),
    "uom": (
        "Unit",
        False,
        (
            "uom", "unit", "units", "unit of measure", "unit of measurement", "per", "measure",
            "um", "u m", "unit of qty", "qty unit", "unit type", "uqc",
        ),
    ),
    "price": (
        "Expected Price",
        False,
        (
            "price", "rate", "expected price", "unit price", "unit rate", "rate per unit",
            "rate unit", "target price", "target rate", "estimated rate", "est rate",
            "estimated price", "budget price", "budget rate", "basic rate", "last purchase price",
            "lpp", "last price", "previous rate", "approx rate", "approx price", "rate rs",
            "rate inr", "unit cost", "cost",
        ),
    ),
    "item_code": (
        "Item Code",
        False,
        (
            "item code", "code", "part no", "part number", "material code", "product code",
            "sku", "catalogue no", "catalog no", "cat no", "model no", "model", "article no",
            "item no", "material no", "drawing no", "ref no", "item id",
        ),
    ),
    "make": (
        "Make / Brand",
        False,
        ("make", "brand", "preferred make", "make brand", "brand name", "manufacturer", "approved make", "oem"),
    ),
    "hsn": ("HSN", False, ("hsn", "hsn code", "hsn sac", "sac", "hsn no")),
    "remarks": (
        "Remarks",
        False,
        (
            "remarks", "remark", "notes", "note", "comment", "comments", "additional info",
            "additional information", "instructions", "special instructions",
        ),
    ),
}

#: Headers that are never a field — skipped outright so a stray "No" in
#: "Sr. No." can't be read as a quantity.
_IGNORED_HEADERS = {
    "s no", "sr no", "sl no", "sno", "srno", "slno", "serial no", "serial number", "s n", "sr",
    "sl", "sn", "no", "line no", "line", "amount", "total", "total amount", "value", "total value",
    "gst", "gst rate", "tax", "igst", "cgst", "sgst", "discount", "net amount", "delivery date",
    "total cost", "total price", "total rate", "amount rs", "amount inr", "gst amount", "tax amount",
}

#: canonical `canonical_fields.code` each field is saved against.
_CANONICAL_CODE = {
    "description": "product_description",
    "quantity": "quantity",
    "uom": "uom",
    "price": "unit_price",
    "item_code": "supplier_sku",
    "make": "brand",
    "hsn": "hsn_code",
    "remarks": "remarks",
}

#: Unit spellings seen on Indian RFQs -> a system unit code. Only true
#: equivalents: a carton is not a box and a roll is not a metre.
_UNIT_ALIASES: dict[str, tuple[str, ...]] = {
    "nos": (
        "nos", "no", "number", "numbers", "nr", "pc", "pcs", "piece", "pieces", "each", "ea",
        "unit", "units", "qty", "item", "items", "ea.", "n",
    ),
    "kg": ("kg", "kgs", "kilogram", "kilograms", "kilo", "kilos"),
    "litre": ("litre", "litres", "liter", "liters", "ltr", "ltrs", "lt", "lts", "l"),
    "metre": ("metre", "metres", "meter", "meters", "mtr", "mtrs", "mt", "m", "rmt", "running metre", "running meter"),
    "box": ("box", "boxes", "bx"),
    "pack": ("pack", "packs", "pk"),
    "pkt": ("pkt", "pkts", "packet", "packets"),
    "set": ("set", "sets"),
    "bag": ("bag", "bags"),
}

_TOTAL_ROW = re.compile(r"^(grand\s*)?(sub\s*)?total\b|^net\s+total\b", re.IGNORECASE)
# The sign is accepted so "-5" is reported as "must be more than zero", not "not a number".
_NUMBER_WITH_UNIT = re.compile(r"^\s*(-?(?:[0-9][0-9,]*(?:\.[0-9]+)?|\.[0-9]+))\s*([A-Za-z][A-Za-z .]*)?\s*$")
_BRACKET = re.compile(r"[\(\[]([^\)\]]*)[\)\]]")


def header_key(value: object) -> str:
    """Header text reduced to comparable words: `"Qty. (Nos)*"` -> `"qty"`,
    `"S.No"` -> `"s no"`. Bracketed notes are dropped here and read
    separately by `_bracket_note`."""
    out = _BRACKET.sub(" ", str(value or "")).lower()
    out = out.replace("&", " and ")
    out = re.sub(r"[^a-z0-9]+", " ", out)
    return re.sub(r"\s+", " ", out).strip()


def _bracket_note(value: object) -> str:
    match = _BRACKET.search(str(value or ""))
    return match.group(1).strip() if match else ""


def _score(key: str, synonyms: tuple[str, ...]) -> tuple[float, str]:
    if not key:
        return 0.0, ""
    if key in synonyms:
        return 1.0, "exact"
    words = f" {key} "
    best = 0.0
    for syn in synonyms:
        if len(syn) < 3:
            continue
        if f" {syn} " in words:
            # "material description" contains "description". Longer
            # synonyms are stronger evidence than a bare "item".
            best = max(best, 0.7 + min(len(syn), 20) / 100)
    if best:
        return best, "contains"
    for syn in synonyms:
        if len(syn) >= 5 and SequenceMatcher(None, key, syn).ratio() >= 0.86:
            return 0.68, "similar"
    return 0.0, ""


def detect_columns(headers: list[str]) -> dict[str, tuple[int, float, str]]:
    """field -> (column index, confidence 0-1, how). Best matches claim
    their column first, so "Unit Price" goes to price (exact) before uom
    can claim it on the word "unit"."""
    keys = [header_key(h) for h in headers]
    scored: list[tuple[float, int, str, str]] = []
    for idx, key in enumerate(keys):
        if not key or key in _IGNORED_HEADERS:
            continue
        for fld, (_label, _req, synonyms) in FIELDS.items():
            score, how = _score(key, synonyms)
            if score:
                scored.append((score, idx, fld, how))
    found: dict[str, tuple[int, float, str]] = {}
    used: set[int] = set()
    for score, idx, fld, how in sorted(scored, key=lambda s: (-s[0], s[1])):
        if fld in found or idx in used:
            continue
        found[fld] = (idx, score, how)
        used.add(idx)
    return found


def signature(headers: list[str]) -> str:
    joined = "|".join(header_key(h) for h in headers)
    return hashlib.sha256(("rfq:" + joined).encode()).hexdigest()


# ------------------------------------------------------------ file reading


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    return str(value).strip()


def read_sheets(filename: str, content: bytes) -> list[tuple[str, list[list[str]]]]:
    """Every sheet as a grid of strings, trailing empties trimmed."""
    name = (filename or "").lower()
    if not content:
        raise ApiError(status.HTTP_400_BAD_REQUEST, CODE_VALIDATION, "The file is empty")
    try:
        if name.endswith(".csv") or name.endswith(".txt"):
            sheets = [("CSV", _read_csv(content))]
        elif name.endswith(".xlsx") or name.endswith(".xlsm"):
            sheets = _read_xlsx(content)
        elif name.endswith(".xls"):
            sheets = _read_xls(content)
        else:
            raise ApiError(
                status.HTTP_400_BAD_REQUEST,
                CODE_VALIDATION,
                "Only Excel (.xlsx, .xls) and .csv files are supported",
            )
    except ApiError:
        raise
    except Exception:
        raise ApiError(
            status.HTTP_400_BAD_REQUEST,
            CODE_VALIDATION,
            "This file could not be read. Open it in Excel and save it again as .xlsx, then retry.",
        )
    out = []
    for sheet_name, grid in sheets:
        grid = [[_cell_text(c) for c in row] for row in grid[:_MAX_ROWS]]
        for row in grid:
            while row and not row[-1]:
                row.pop()
        out.append((sheet_name, grid))
    return out


def _read_csv(content: bytes) -> list[list]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            decoded = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        decoded = content.decode("utf-8", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(decoded), dialect)]


def _read_xlsx(content: bytes) -> list[tuple[str, list[list]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        return [
            (ws.title, [list(r) for r in ws.iter_rows(max_row=_MAX_ROWS, values_only=True)])
            for ws in workbook.worksheets
        ]
    finally:
        workbook.close()


def _read_xls(content: bytes) -> list[tuple[str, list[list]]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheets = []
    for sheet in book.sheets():
        grid = []
        for r in range(min(sheet.nrows, _MAX_ROWS)):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            grid.append(row)
        sheets.append((sheet.name, grid))
    return sheets


# ----------------------------------------------------------- the pipeline


@dataclass
class Layout:
    sheet_index: int
    header_row: int  # 0-based within the sheet
    columns: dict[str, int]
    confidence: dict[str, tuple[float, str]] = field(default_factory=dict)
    source: str = "auto"  # auto | saved | manual
    mapping_id: Optional[UUID] = None
    #: field -> why it was mapped, for fields column_assist filled in.
    reasons: dict[str, str] = field(default_factory=dict)


def _auto_layout(sheets: list[tuple[str, list[list[str]]]]) -> Optional[Layout]:
    best: Optional[tuple[float, Layout]] = None
    for s_idx, (_name, grid) in enumerate(sheets):
        for r_idx, row in enumerate(grid[:_HEADER_SCAN_ROWS]):
            if sum(1 for c in row if c) < 2:
                continue
            found = detect_columns(row)
            if "description" not in found and "quantity" not in found:
                continue
            score = sum(conf for _i, conf, _h in found.values())
            score += 2.0 * ("description" in found) + 2.0 * ("quantity" in found)
            # Prefer a header with data under it over a lone title line.
            score += min(sum(1 for r in grid[r_idx + 1 : r_idx + 6] if any(r)), 5) * 0.05
            if best is None or score > best[0]:
                best = (
                    score,
                    Layout(
                        sheet_index=s_idx,
                        header_row=r_idx,
                        columns={f: v[0] for f, v in found.items()},
                        confidence={f: (v[1], v[2]) for f, v in found.items()},
                    ),
                )
    return best[1] if best else None


async def _saved_layouts(session: AsyncSession, company_id: str) -> dict[str, tuple[UUID, dict[str, int]]]:
    rows = (
        await session.execute(
            text(
                "SELECT m.id, m.column_signature, cf.code, f.source_column_index "
                "FROM document_schema_mappings m "
                "JOIN document_schema_mapping_fields f ON f.mapping_id = m.id "
                "JOIN canonical_fields cf ON cf.id = f.canonical_field_id "
                "WHERE m.company_id = :c AND m.document_type = :t AND m.status = 'confirmed' "
                "ORDER BY m.version DESC"
            ),
            {"c": company_id, "t": DOCUMENT_TYPE},
        )
    ).mappings().all()
    by_code = {v: k for k, v in _CANONICAL_CODE.items()}
    out: dict[str, tuple[UUID, dict[str, int]]] = {}
    for r in rows:
        sig = r["column_signature"]
        if sig in out and out[sig][0] != r["id"]:
            continue  # an older version of the same layout
        entry = out.setdefault(sig, (r["id"], {}))
        fld = by_code.get(r["code"])
        if fld and r["source_column_index"] is not None:
            entry[1][fld] = r["source_column_index"]
    return out


def _saved_layout(sheets, saved) -> Optional[Layout]:
    if not saved:
        return None
    for s_idx, (_name, grid) in enumerate(sheets):
        for r_idx, row in enumerate(grid[:_HEADER_SCAN_ROWS]):
            if sum(1 for c in row if c) < 2:
                continue
            hit = saved.get(signature(row))
            if hit:
                return Layout(
                    sheet_index=s_idx,
                    header_row=r_idx,
                    columns=dict(hit[1]),
                    confidence={f: (1.0, "saved") for f in hit[1]},
                    source="saved",
                    mapping_id=hit[0],
                )
    return None


class UnitResolver:
    def __init__(self, uoms: list[tuple[UUID, str, str]]):
        self.by_key: dict[str, tuple[UUID, str]] = {}
        for uid, code, name in uoms:
            for key in (code, name):
                self.by_key.setdefault(str(key).strip().lower(), (uid, code))
        for code, aliases in _UNIT_ALIASES.items():
            target = self.by_key.get(code)
            if target:
                for alias in aliases:
                    self.by_key.setdefault(alias, target)

    def resolve(self, raw: str) -> Optional[tuple[UUID, str]]:
        key = re.sub(r"[^a-z0-9 ]+", " ", (raw or "").lower())
        key = re.sub(r"\s+", " ", key).strip()
        if not key:
            return None
        return self.by_key.get(key) or (self.by_key.get(key[:-1]) if key.endswith("s") else None)


def _parse_number(raw: str) -> Optional[float]:
    cleaned = re.sub(r"(?i)rs\.?|inr|₹|/-", "", raw or "").replace(",", "").strip()
    if not cleaned:
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def _split_quantity(raw: str) -> tuple[Optional[float], str]:
    """`"1,200 Nos"` -> (1200.0, "Nos"); `"abc"` -> (None, "")."""
    match = _NUMBER_WITH_UNIT.match(raw or "")
    if not match:
        return None, ""
    return float(match.group(1).replace(",", "")), (match.group(2) or "").strip().rstrip(".")


@dataclass
class ParseResult:
    rows: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    sheet_names: list[str] = field(default_factory=list)
    sheet_index: int = 0
    header_row: Optional[int] = None  # 1-based, as a person counts rows
    headers: list[str] = field(default_factory=list)
    sample_rows: list[list[str]] = field(default_factory=list)
    columns: list[dict] = field(default_factory=list)
    missing_required: list[str] = field(default_factory=list)
    layout_source: str = "auto"
    signature: Optional[str] = None
    column_map: dict[str, int] = field(default_factory=dict)

    def as_dict(self) -> dict:
        return {
            "rows": self.rows,
            "errors": self.errors,
            "warnings": self.warnings,
            "row_count": len(self.rows),
            "sheet_names": self.sheet_names,
            "sheet_index": self.sheet_index,
            "header_row": self.header_row,
            "headers": self.headers,
            "sample_rows": self.sample_rows,
            "columns": self.columns,
            "missing_required": self.missing_required,
            "layout_source": self.layout_source,
            "column_map": self.column_map,
        }


def _widest_row(grid: list[list[str]]) -> int:
    """When no heading is recognisable, the heading row is the first row as
    wide as the table itself — not the letterhead line above it."""
    top = grid[:_HEADER_SCAN_ROWS]
    widest = max((sum(1 for c in r if c) for r in top), default=0)
    if widest == 0:
        return 0
    need = max(2, int(widest * 0.6 + 0.5))
    # A single-column file has no row that wide; its first row is the heading.
    return next((i for i, r in enumerate(top) if sum(1 for c in r if c) >= need), 0)


#: RFQ field -> the generic role `column_assist` understands.
_ASSIST_ROLE = {
    "description": "description",
    "quantity": "quantity",
    "uom": "uom",
    "price": "price",
    "item_code": "item_code",
    "make": "make",
    "hsn": "hsn",
    "remarks": "remarks",
}


async def _assist(layout: "Layout", headers: list[str], data_rows: list[list[str]], units: "UnitResolver", result: "ParseResult") -> None:
    """Fill fields the synonyms missed from the column values, the
    headings' meaning and — for a required field still missing — a language
    model. Suggestions only; the person confirms them in the mapping step."""
    specs = [
        column_assist.FieldSpec(
            key=f, label=label, role=_ASSIST_ROLE[f], required=required, synonyms=synonyms
        )
        for f, (label, required, synonyms) in FIELDS.items()
        if f not in layout.columns
    ]
    taken = set(layout.columns.values()) | {
        i for i, h in enumerate(headers) if header_key(h) in _IGNORED_HEADERS
    }
    if not specs or len(taken) >= len(headers):
        return
    rows = [r for r in data_rows if any(r)][:30]
    try:
        assisted = await column_assist.suggest(
            headers,
            rows,
            specs,
            taken_columns=taken,
            unit_words=column_assist.DEFAULT_UNIT_WORDS | frozenset(units.by_key),
        )
    except Exception:  # never let an optional helper fail an import
        log.exception("column assist failed")
        return
    for fld, sug in assisted.suggestions.items():
        layout.columns[fld] = sug.column_index
        layout.confidence[fld] = (sug.confidence / 100, sug.source)
        layout.reasons[fld] = sug.reason
    result.warnings.extend(assisted.notes)


def parse_column_map(raw: Optional[str]) -> Optional[dict[str, int]]:
    """The manual mapping as sent by the import screen: a JSON object of
    field -> 0-based column index (or null/-1 for "not in this file")."""
    if not raw:
        return None
    try:
        data = json.loads(raw)
        assert isinstance(data, dict)
    except Exception:
        raise ApiError(status.HTTP_400_BAD_REQUEST, CODE_VALIDATION, "column_map must be a JSON object")
    out: dict[str, int] = {}
    for fld, idx in data.items():
        if fld not in FIELDS:
            raise ApiError(status.HTTP_400_BAD_REQUEST, CODE_VALIDATION, f"Unknown field '{fld}' in column_map")
        if idx is None or idx == "" or int(idx) < 0:
            continue
        out[fld] = int(idx)
    if len(set(out.values())) != len(out):
        raise ApiError(status.HTTP_400_BAD_REQUEST, CODE_VALIDATION, "Two fields can't use the same column")
    return out


async def parse(
    session: AsyncSession,
    *,
    company_id: str,
    filename: str,
    content: bytes,
    column_map: Optional[dict[str, int]] = None,
    header_row: Optional[int] = None,
    sheet_index: Optional[int] = None,
    default_uom_id: Optional[UUID] = None,
) -> ParseResult:
    """Read the file and build RFQ lines. `column_map`, `header_row`
    (1-based) and `sheet_index` are the person's overrides from the mapping
    step; anything they don't override is detected."""
    result = ParseResult()
    sheets = read_sheets(filename, content)
    result.sheet_names = [name for name, _g in sheets]
    if not any(any(any(c for c in row) for row in grid) for _n, grid in sheets):
        result.errors.append("The file has no rows")
        return result

    if sheet_index is not None and not 0 <= sheet_index < len(sheets):
        raise ApiError(status.HTTP_400_BAD_REQUEST, CODE_VALIDATION, "That sheet is not in the file")

    # ---- which sheet, which header row, which columns
    layout: Optional[Layout] = None
    if header_row is None and column_map is None:
        layout = _saved_layout(
            sheets if sheet_index is None else [sheets[sheet_index]],
            await _saved_layouts(session, company_id),
        )
        if layout and sheet_index is not None:
            layout.sheet_index = sheet_index
        if layout is None:
            candidates = sheets if sheet_index is None else [sheets[sheet_index]]
            layout = _auto_layout(candidates)
            if layout and sheet_index is not None:
                layout.sheet_index = sheet_index
    if layout is None:
        # Nothing recognisable anywhere: the table is on the fullest sheet,
        # not the cover page in front of it.
        s_idx = sheet_index if sheet_index is not None else max(
            range(len(sheets)), key=lambda i: sum(1 for r in sheets[i][1] if any(r))
        )
        grid = sheets[s_idx][1]
        if header_row is not None:
            r_idx = header_row - 1
            if not 0 <= r_idx < len(grid):
                raise ApiError(status.HTTP_400_BAD_REQUEST, CODE_VALIDATION, "The header row is outside the sheet")
        else:
            auto = _auto_layout([sheets[s_idx]])
            r_idx = auto.header_row if auto else _widest_row(grid)
        found = detect_columns(grid[r_idx]) if r_idx < len(grid) else {}
        layout = Layout(
            sheet_index=s_idx,
            header_row=r_idx,
            columns={f: v[0] for f, v in found.items()},
            confidence={f: (v[1], v[2]) for f, v in found.items()},
        )
    if column_map is not None:
        layout.columns = dict(column_map)
        layout.source = "manual"
        layout.confidence = {f: (1.0, "manual") for f in column_map}

    uom_rows = (
        await session.execute(
            text(
                "SELECT id, code, name FROM uoms WHERE (company_id IS NULL OR company_id = :c) AND is_active"
            ),
            {"c": company_id},
        )
    ).all()
    units = UnitResolver([(r[0], r[1], r[2]) for r in uom_rows])

    sheet_name, grid = sheets[layout.sheet_index]
    headers = grid[layout.header_row] if layout.header_row < len(grid) else []
    # Signed before padding, exactly as `_saved_layout` signs a raw row.
    result.signature = signature(headers)
    width = max([len(headers)] + [len(r) for r in grid[layout.header_row + 1 : layout.header_row + 50]])
    headers = headers + [""] * (width - len(headers))
    for fld, idx in layout.columns.items():
        if idx >= width:
            raise ApiError(status.HTTP_400_BAD_REQUEST, CODE_VALIDATION, f"Column {idx + 1} is not in the file")

    data_rows = grid[layout.header_row + 1 :]
    if layout.source == "auto":
        await _assist(layout, headers, data_rows, units, result)
    result.sheet_index = layout.sheet_index
    result.header_row = layout.header_row + 1
    result.headers = [h or f"Column {i + 1}" for i, h in enumerate(headers)]
    result.sample_rows = [(r + [""] * width)[:width] for r in data_rows if any(r)][:5]
    result.layout_source = layout.source
    result.column_map = dict(layout.columns)
    result.columns = [
        {
            "field": fld,
            "label": label,
            "required": required,
            "column_index": layout.columns.get(fld),
            "header": result.headers[layout.columns[fld]] if fld in layout.columns else None,
            "confidence": round(layout.confidence.get(fld, (0.0, ""))[0] * 100) if fld in layout.columns else 0,
            "match": layout.confidence.get(fld, (0.0, ""))[1] if fld in layout.columns else None,
            "reason": layout.reasons.get(fld) if fld in layout.columns else None,
        }
        for fld, (label, required, _syn) in FIELDS.items()
    ]
    result.missing_required = [f for f, (_l, req, _s) in FIELDS.items() if req and f not in layout.columns]
    if "description" in result.missing_required and "item_code" in layout.columns:
        result.missing_required.remove("description")  # the code alone can describe the line
    if result.missing_required:
        names = ", ".join(FIELDS[f][0] for f in result.missing_required)
        result.errors.append(
            f"Couldn't tell which column holds: {names}. Pick it in the column mapping below."
        )
        return result

    # ---- units
    default_unit: Optional[tuple[UUID, str]] = None
    if default_uom_id is not None:
        default_unit = next(((r[0], r[1]) for r in uom_rows if r[0] == default_uom_id), None)
        if default_unit is None:
            raise ApiError(status.HTTP_400_BAD_REQUEST, CODE_VALIDATION, "The default unit is not available")
    # "Qty (Nos)" / "Unit (Kg)" — a unit written once, in the header.
    header_unit = None
    for fld in ("quantity", "uom"):
        idx = layout.columns.get(fld)
        if idx is not None:
            header_unit = header_unit or units.resolve(_bracket_note(headers[idx]))

    def cell(raw: list[str], fld: str) -> str:
        idx = layout.columns.get(fld)
        return raw[idx].strip() if idx is not None and idx < len(raw) and raw[idx] else ""

    no_unit_rows: list[int] = []
    unknown_units: dict[str, list[int]] = {}
    defaulted: dict[str, list[int]] = {}
    headings: list[str] = []

    for offset, raw in enumerate(data_rows):
        line = layout.header_row + offset + 2  # 1-based sheet row
        if not any(raw):
            continue
        description = cell(raw, "description")
        code = cell(raw, "item_code")
        qty_raw = cell(raw, "quantity")
        if not description and code:
            description = code
        if not description:
            continue
        if _TOTAL_ROW.match(description):
            continue
        if not qty_raw:
            # "ELECTRICAL ITEMS" between the lines — a heading, not an item.
            headings.append(f"row {line} ('{description[:40]}')")
            continue

        quantity, qty_unit = _split_quantity(qty_raw)
        if quantity is None:
            result.errors.append(f"Row {line}: quantity '{qty_raw}' is not a number")
            continue
        if quantity <= 0:
            result.errors.append(f"Row {line}: quantity must be more than zero")
            continue

        unit_text = cell(raw, "uom") or qty_unit
        unit = units.resolve(unit_text) if unit_text else header_unit
        if unit is None:
            if default_unit is not None:
                unit = default_unit
                if unit_text:
                    defaulted.setdefault(unit_text, []).append(line)
            elif unit_text:
                unknown_units.setdefault(unit_text, []).append(line)
                continue
            else:
                no_unit_rows.append(line)
                continue

        price_raw = cell(raw, "price")
        price = _parse_number(price_raw) if price_raw else 0.0
        if price is None or price < 0:
            result.warnings.append(f"Row {line}: price '{price_raw}' isn't a number — left blank")
            price = 0.0

        notes = []
        if code and code != description:
            notes.append(f"Code: {code}")
        if cell(raw, "make"):
            notes.append(f"Make: {cell(raw, 'make')}")
        if cell(raw, "hsn"):
            notes.append(f"HSN: {cell(raw, 'hsn')}")
        if cell(raw, "remarks"):
            notes.append(cell(raw, "remarks"))

        result.rows.append(
            {
                "source_row": line,
                "description": description,
                "quantity": quantity,
                "uom_id": unit[0],
                "uom_code": unit[1],
                "expected_price": price,
                "remarks": " · ".join(notes) or None,
            }
        )

    def _rows(lines: list[int]) -> str:
        shown = ", ".join(str(n) for n in lines[:6])
        return f"row {shown}" if len(lines) == 1 else f"rows {shown}{'…' if len(lines) > 6 else ''}"

    for unit_text, lines in unknown_units.items():
        result.errors.append(
            f"Unit '{unit_text}' ({_rows(lines)}) isn't set up. Choose a default unit, or add '{unit_text}' "
            "under Settings → Units."
        )
    if no_unit_rows:
        result.errors.append(
            f"No unit given on {_rows(no_unit_rows)}. Choose a default unit, or map the unit column."
        )
    for unit_text, lines in defaulted.items():
        result.warnings.append(
            f"Unit '{unit_text}' isn't set up — {_rows(lines)} use the default unit {default_unit[1]}."
        )
    if headings:
        shown = "; ".join(headings[:4]) + ("…" if len(headings) > 4 else "")
        result.warnings.append(f"Skipped {len(headings)} line(s) with no quantity (headings?): {shown}")
    if not result.rows and not result.errors:
        result.errors.append("No item lines found below the header row")
    return result


async def save_layout(
    session: AsyncSession, *, company_id: str, user_id: str, result: ParseResult, filename: str
) -> None:
    """Remember the layout just imported so the next file with the same
    header row maps itself. Importing *is* the confirmation (BR-AI-09): a
    person has seen the preview and pressed Confirm."""
    if not result.signature or not result.column_map:
        return
    codes = {_CANONICAL_CODE[f]: idx for f, idx in result.column_map.items()}
    fields = (
        await session.execute(
            text("SELECT id, code FROM canonical_fields WHERE code = ANY(:codes)"),
            {"codes": list(codes)},
        )
    ).all()
    field_ids = {code: fid for fid, code in fields}

    existing = (
        await session.execute(
            text(
                "SELECT id FROM document_schema_mappings WHERE company_id = :c AND document_type = :t "
                "AND column_signature = :s AND supplier_id IS NULL AND status = 'confirmed' "
                "ORDER BY version DESC LIMIT 1"
            ),
            {"c": company_id, "t": DOCUMENT_TYPE, "s": result.signature},
        )
    ).scalar_one_or_none()
    if existing:
        mapping_id = existing
        await session.execute(
            text(
                "UPDATE document_schema_mappings SET usage_count = usage_count + 1, last_used_at = now(), "
                "header_row_index = :hdr, data_start_row = :start WHERE id = :id"
            ),
            {"id": mapping_id, "hdr": result.header_row - 1, "start": result.header_row},
        )
        await session.execute(text("DELETE FROM document_schema_mapping_fields WHERE mapping_id = :m"), {"m": mapping_id})
    else:
        version = (
            await session.execute(
                text(
                    "SELECT COALESCE(MAX(version), 0) + 1 FROM document_schema_mappings "
                    "WHERE company_id = :c AND supplier_id IS NULL AND document_type = :t AND column_signature = :s"
                ),
                {"c": company_id, "t": DOCUMENT_TYPE, "s": result.signature},
            )
        ).scalar_one()
        mapping_id = uuid4()
        name = filename.lower()
        await session.execute(
            text(
                "INSERT INTO document_schema_mappings (id, company_id, supplier_id, document_type, file_format, "
                " mapping_name, column_signature, header_row_index, data_start_row, sheet_name, status, "
                " confidence, confirmed_by, confirmed_at, usage_count, last_used_at, version) "
                "VALUES (:id, :c, NULL, :t, :fmt, :name, :sig, :hdr, :start, :sheet, 'confirmed', 100, :u, now(), "
                " 1, now(), :ver)"
            ),
            {
                "id": mapping_id,
                "c": company_id,
                "t": DOCUMENT_TYPE,
                "fmt": "csv" if name.endswith(".csv") else "xlsx",
                "name": f"RFQ layout — {filename}"[:200],
                "sig": result.signature,
                "hdr": result.header_row - 1,
                "start": result.header_row,
                "sheet": result.sheet_names[result.sheet_index] if result.sheet_names else None,
                "u": user_id,
                "ver": version,
            },
        )
    seen_headers: set[str] = set()
    for code, idx in codes.items():
        if code not in field_ids:
            continue  # vocabulary row not present in this database yet
        source_column = result.headers[idx] if idx < len(result.headers) else f"Column {idx + 1}"
        if source_column in seen_headers:  # two columns both headed "Rate"
            source_column = f"{source_column} (column {idx + 1})"
        seen_headers.add(source_column)
        await session.execute(
            text(
                "INSERT INTO document_schema_mapping_fields (id, mapping_id, source_column, source_column_index, "
                " canonical_field_id, transform, is_required, confidence, confirmed_by, confirmed_at) "
                "VALUES (:id, :m, :col, :idx, :cf, 'none', :req, 100, :u, now())"
            ),
            {
                "id": uuid4(),
                "m": mapping_id,
                "col": source_column,
                "idx": idx,
                "cf": field_ids[code],
                "req": code in ("product_description", "quantity"),
                "u": user_id,
            },
        )
